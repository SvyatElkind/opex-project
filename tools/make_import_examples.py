# -*- coding: utf-8 -*-
"""Generate the CSV/XLSX import example files shipped to users.

Run from the repo root:  python make_examples.py
Output: opex_tool_frontend/public/examples/
"""

import csv
import io
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_DIR = os.path.join('opex_tool_frontend', 'public', 'examples')

# --- Canonical column order (union of GV + DOK columns) ----------------------

COLUMNS = [
    'TIPS', 'SAITE',
    # GV
    'SĒRIJAS_KODS', 'NOSAUKUMS', 'DATUMS_NO', 'DATUMS_LĪDZ', 'DATUMA_PRECIZITĀTE',
    'DATUMA_PIEZĪMES', 'VALODA', 'SATURS', 'PIEZĪMES', 'SISTEMATIZĀCIJA',
    'APJOMS', 'APJOMA_MĒRVIENĪBA',
    'PIEEJAMĪBA', 'PIEEJAMĪBAS_PAMATOJUMS', 'SLEPENĪBA', 'SLEPENĪBAS_PIEZĪMES',
    'KOPIJA', 'ARHĪVA_VĒSTURE',
    # DOK
    'DATUMS', 'REĢ_NR', 'IZVEIDOŠANAS_DATUMS', 'NOSŪTĪŠANAS_DATUMS',
    'NOSŪTĪTĀJA_REĢ_NR', 'LIETAS_NR', 'GRUPA', 'ATSLĒGVĀRDI', 'ANOTĀCIJA',
    'TEHNISKĀ_INFORMĀCIJA', 'IEROBEŽOJUMA_DATUMS', 'IEROBEŽOJUMA_PIEZĪMES',
    'LIETOŠANAS_NOSACĪJUMI',
]


def row(**kwargs):
    """Build a full-width row dict; unspecified columns stay empty."""
    unknown = set(kwargs) - set(COLUMNS)
    assert not unknown, f'unknown columns: {unknown}'
    return {column: kwargs.get(column, '') for column in COLUMNS}


# --- Mixed example: items with their records (usage C) -----------------------

MIXED_ROWS = [
    row(TIPS='GV', SAITE='A', SĒRIJAS_KODS='1.2',
        NOSAUKUMS='Domes sēžu protokoli 2020. gada I ceturksnis',
        DATUMS_NO='01.01.2020', DATUMS_LĪDZ='31.03.2020', DATUMA_PRECIZITĀTE='diena',
        VALODA='Latviešu',
        SATURS='Domes sēžu protokoli un tiem pievienotie lēmumu projekti.',
        PIEZĪMES='Pārņemts no lietu nomenklatūras 2020.',
        SISTEMATIZĀCIJA='Hronoloģiski', PIEEJAMĪBA='Vispārēja', SLEPENĪBA='Publisks'),
    row(TIPS='DOK', SAITE='A', NOSAUKUMS='Domes sēdes protokols Nr. 1',
        DATUMS='15.01.2020', REĢ_NR='1-15/1',
        IZVEIDOŠANAS_DATUMS='15.01.2020', NOSŪTĪŠANAS_DATUMS='16.01.2020',
        LIETAS_NR='1-15', GRUPA='Iekšējs', VALODA='Latviešu',
        ATSLĒGVĀRDI='protokols,dome,sēde',
        ANOTĀCIJA='Kārtējā domes sēde. Izskatīti 12 jautājumi.',
        PIEEJAMĪBA='Vispārēja'),
    row(TIPS='DOK', SAITE='A', NOSAUKUMS='Domes sēdes protokols Nr. 2',
        DATUMS='20.02.2020', REĢ_NR='1-15/2',
        IZVEIDOŠANAS_DATUMS='20.02.2020', NOSŪTĪŠANAS_DATUMS='21.02.2020',
        LIETAS_NR='1-15', GRUPA='Iekšējs', VALODA='Latviešu',
        ATSLĒGVĀRDI='protokols,dome,sēde',
        ANOTĀCIJA='Ārkārtas sēde par budžeta grozījumiem.',
        PIEEJAMĪBA='Vispārēja'),
    row(TIPS='DOK', SAITE='A', NOSAUKUMS='Domes sēdes protokols Nr. 3',
        DATUMS='18.03.2020', REĢ_NR='1-15/3',
        IZVEIDOŠANAS_DATUMS='18.03.2020', NOSŪTĪŠANAS_DATUMS='19.03.2020',
        LIETAS_NR='1-15', GRUPA='Iekšējs', VALODA='Latviešu',
        ANOTĀCIJA='Sēde attālināti.', PIEEJAMĪBA='Vispārēja'),

    row(TIPS='GV', SAITE='B', SĒRIJAS_KODS='1.3',
        NOSAUKUMS='Sarakste ar VARAM 2020. gads',
        DATUMS_NO='01.01.2020', DATUMS_LĪDZ='31.12.2020', DATUMA_PRECIZITĀTE='diena',
        VALODA='Latviešu; Angļu',
        SATURS='Iestādes sarakste ar Vides aizsardzības un reģionālās attīstības ministriju.',
        PIEEJAMĪBA='Vispārēja', SLEPENĪBA='Iekšējs',
        SLEPENĪBAS_PIEZĪMES='Dienesta vajadzībām'),
    row(TIPS='DOK', SAITE='B', NOSAUKUMS='Vēstule VARAM par teritorijas plānojumu',
        DATUMS='05.03.2020', REĢ_NR='1-19/7',
        IZVEIDOŠANAS_DATUMS='04.03.2020', NOSŪTĪŠANAS_DATUMS='05.03.2020',
        NOSŪTĪTĀJA_REĢ_NR='4.1-2/1234', LIETAS_NR='1-19', GRUPA='Nosūtīts',
        VALODA='Latviešu', ATSLĒGVĀRDI='teritorijas plānojums,VARAM',
        ANOTĀCIJA='Atbilde uz ministrijas 2020. gada 20. februāra pieprasījumu.',
        PIEEJAMĪBA='Vispārēja'),
    row(TIPS='DOK', SAITE='B', NOSAUKUMS='VARAM atbilde par teritorijas plānojumu',
        DATUMS='30.03.2020', REĢ_NR='1-19/8',
        IZVEIDOŠANAS_DATUMS='27.03.2020', NOSŪTĪŠANAS_DATUMS='30.03.2020',
        NOSŪTĪTĀJA_REĢ_NR='4.1-2/1300', LIETAS_NR='1-19', GRUPA='Saņemts',
        VALODA='Latviešu', PIEEJAMĪBA='Vispārēja'),

    row(TIPS='GV', SAITE='C', SĒRIJAS_KODS='2.1',
        NOSAUKUMS='Personāla lietas 2020. gads',
        DATUMS_NO='01.01.2020', DATUMS_LĪDZ='31.12.2020', DATUMA_PRECIZITĀTE='diena',
        VALODA='Latviešu',
        SATURS='Darbinieku personāla lietas.',
        PIEEJAMĪBA='Sensitīvi dati',
        PIEEJAMĪBAS_PAMATOJUMS='Satur fizisko personu datus; VDAR 9. pants',
        SLEPENĪBA='Konfidenciāls', SLEPENĪBAS_PIEZĪMES='Piekļuve tikai personāla nodaļai'),
    row(TIPS='DOK', SAITE='C', NOSAUKUMS='Rīkojums par darbinieka pieņemšanu darbā',
        DATUMS='10.02.2020', REĢ_NR='3-1/15',
        IZVEIDOŠANAS_DATUMS='10.02.2020', NOSŪTĪŠANAS_DATUMS='10.02.2020',
        LIETAS_NR='3-1', GRUPA='Iekšējs', VALODA='Latviešu',
        PIEEJAMĪBA='Ierobežota', IEROBEŽOJUMA_DATUMS='31.12.2100',
        IEROBEŽOJUMA_PIEZĪMES='Personas dati',
        LIETOŠANAS_NOSACĪJUMI='Tikai ar atļauju'),
]

# --- Items-only example (usage A) -------------------------------------------

ITEMS_ONLY_ROWS = [
    row(TIPS='GV', SĒRIJAS_KODS='1.2',
        NOSAUKUMS='Domes sēžu protokoli 2020. gada I ceturksnis',
        DATUMS_NO='01.01.2020', DATUMS_LĪDZ='31.03.2020',
        VALODA='Latviešu', PIEEJAMĪBA='Vispārēja', SLEPENĪBA='Publisks'),
    row(TIPS='GV', SĒRIJAS_KODS='1.2',
        NOSAUKUMS='Domes sēžu protokoli 2020. gada II ceturksnis',
        DATUMS_NO='01.04.2020', DATUMS_LĪDZ='30.06.2020',
        VALODA='Latviešu', PIEEJAMĪBA='Vispārēja', SLEPENĪBA='Publisks'),
    row(TIPS='GV', SĒRIJAS_KODS='1.2',
        NOSAUKUMS='Domes sēžu protokoli 2020. gada III ceturksnis',
        DATUMS_NO='01.07.2020', DATUMS_LĪDZ='30.09.2020',
        VALODA='Latviešu', PIEEJAMĪBA='Vispārēja', SLEPENĪBA='Publisks'),
    row(TIPS='GV', SĒRIJAS_KODS='1.2',
        NOSAUKUMS='Domes sēžu protokoli 2020. gada IV ceturksnis',
        DATUMS_NO='01.10.2020', DATUMS_LĪDZ='31.12.2020',
        VALODA='Latviešu', PIEEJAMĪBA='Vispārēja', SLEPENĪBA='Publisks'),
    row(TIPS='GV', SĒRIJAS_KODS='1.4',
        NOSAUKUMS='Gada pārskats 2020',
        DATUMS_NO='2020', DATUMS_LĪDZ='2020', DATUMA_PRECIZITĀTE='gads',
        VALODA='Latviešu',
        DATUMA_PIEZĪMES='Precīzs sagatavošanas datums nav zināms',
        PIEEJAMĪBA='Vispārēja', SLEPENĪBA='Publisks'),
]

# --- Records-only example (usage B) -----------------------------------------

RECORDS_ONLY_ROWS = [
    row(TIPS='DOK', SAITE='GV:12', NOSAUKUMS='Domes sēdes protokols Nr. 1',
        DATUMS='15.01.2020', REĢ_NR='1-15/1',
        IZVEIDOŠANAS_DATUMS='15.01.2020', NOSŪTĪŠANAS_DATUMS='16.01.2020',
        LIETAS_NR='1-15', GRUPA='Iekšējs', VALODA='Latviešu',
        PIEEJAMĪBA='Vispārēja'),
    row(TIPS='DOK', SAITE='GV:12', NOSAUKUMS='Domes sēdes protokols Nr. 2',
        DATUMS='20.02.2020', REĢ_NR='1-15/2',
        IZVEIDOŠANAS_DATUMS='20.02.2020', NOSŪTĪŠANAS_DATUMS='21.02.2020',
        LIETAS_NR='1-15', GRUPA='Iekšējs', VALODA='Latviešu',
        PIEEJAMĪBA='Vispārēja'),
    row(TIPS='DOK', SAITE='GV:13', NOSAUKUMS='Vēstule VARAM par teritorijas plānojumu',
        DATUMS='05.03.2020', REĢ_NR='1-19/7',
        IZVEIDOŠANAS_DATUMS='04.03.2020', NOSŪTĪŠANAS_DATUMS='05.03.2020',
        LIETAS_NR='1-19', GRUPA='Nosūtīts', VALODA='Latviešu',
        PIEEJAMĪBA='Vispārēja'),
]


def write_csv(path, rows):
    """Write UTF-8 **with BOM**, ';' delimiter, CRLF — what Excel on a Latvian
    Windows opens correctly by double-click."""
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=COLUMNS, delimiter=';',
                           lineterminator='\r\n', quoting=csv.QUOTE_MINIMAL)
    writer.writeheader()
    for r in rows:
        writer.writerow(r)
    with io.open(path, 'w', encoding='utf-8-sig', newline='') as handle:
        handle.write(buffer.getvalue())
    print(f'  {path}  ({len(rows)} rindas)')


# --- XLSX -------------------------------------------------------------------

HEADER_FILL = PatternFill('solid', fgColor='2F5061')
GV_FILL = PatternFill('solid', fgColor='EAF2F8')
DOK_FILL = PatternFill('solid', fgColor='FDF6E3')
TITLE_FONT = Font(bold=True, size=13)
HEADER_FONT = Font(bold=True, color='FFFFFF', size=10)
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Which columns belong to which row type — used by the INSTRUKCIJA sheet.
GV_COLUMNS = COLUMNS[2:20]
DOK_COLUMNS = COLUMNS[20:]

REQUIRED_GV = {'SĒRIJAS_KODS', 'NOSAUKUMS', 'DATUMS_NO', 'DATUMS_LĪDZ', 'VALODA'}
REQUIRED_DOK = {'NOSAUKUMS', 'DATUMS', 'REĢ_NR', 'IZVEIDOŠANAS_DATUMS',
                'NOSŪTĪŠANAS_DATUMS', 'LIETAS_NR', 'VALODA'}

COLUMN_DOCS = {
    'TIPS': ('GV vai DOK', 'GV = glabājamā vienība, DOK = dokuments. Obligāta katrai rindai.'),
    'SAITE': ('brīva atslēga / GV:<nr>', 'Saista DOK rindu ar vienību. Tukšs = tuvākā augstāk esošā GV rinda. "GV:12" = jau esoša vienība ar GV numuru 12.'),
    'SĒRIJAS_KODS': ('1 / 1.2 / 1.2.3', 'Bez sākuma nullēm. Obligāts GV rindām.'),
    'NOSAUKUMS': ('teksts', 'GV rindā — vienības nosaukums (līdz 1000 z.), DOK rindā — dokumenta nosaukums (līdz 500 z.). Obligāts.'),
    'DATUMS_NO': ('01.01.2020 / 2020-01-01 / 2020', 'Vienības sākuma datums. Obligāts GV rindām.'),
    'DATUMS_LĪDZ': ('31.03.2020 / 2020', 'Vienības beigu datums. Nedrīkst būt vēlāks par uzskaites saraksta beigu datumu. Obligāts GV rindām.'),
    'DATUMA_PRECIZITĀTE': ('diena / mēnesis / gads', 'Ja tukšs — nosaka pēc datuma pieraksta ("2020" = gads).'),
    'DATUMA_PIEZĪMES': ('teksts', 'Piezīme par datējumu, piem. ja precīzs datums nav zināms.'),
    'VALODA': ('Latviešu; Krievu', 'Vairākas valodas atdala ar semikolu. Obligāta (izņemot Foto sarakstus).'),
    'SATURS': ('teksts', 'Vienības satura apraksts (līdz 2000 z.). Obligāts Foto/Video/Skaņas sarakstos.'),
    'PIEZĪMES': ('teksts', 'Piezīmes (GV — līdz 1000 z., DOK — līdz 500 z.).'),
    'SISTEMATIZĀCIJA': ('teksts', 'Kā dokumenti vienības iekšienē ir sakārtoti, piem. "Hronoloģiski".'),
    'APJOMS': ('skaitlis', 'Tikai papīra (neelektroniskiem) sarakstiem.'),
    'APJOMA_MĒRVIENĪBA': ('Lapas / Dokumenti / Glabājamās vienības', 'Tikai papīra sarakstiem. Noklusējums — Lapas.'),
    'PIEEJAMĪBA': ('GV: Vispārēja/Ierobežota/Sensitīvi dati; DOK: Vispārēja/Ierobežota', 'Noklusējums — Vispārēja.'),
    'PIEEJAMĪBAS_PAMATOJUMS': ('teksts', 'Obligāts GV rindai, ja PIEEJAMĪBA nav "Vispārēja".'),
    'SLEPENĪBA': ('Publisks / Iekšējs / Konfidenciāls / Slepens', 'Tikai GV rindām. Noklusējums — Publisks.'),
    'SLEPENĪBAS_PIEZĪMES': ('teksts', 'Tikai GV rindām.'),
    'KOPIJA': ('teksts', 'Ziņas par kopijām. Tikai GV rindām.'),
    'ARHĪVA_VĒSTURE': ('teksts', 'Ziņas par dokumentu iepriekšējo glabāšanu. Tikai GV rindām.'),
    'DATUMS': ('15.01.2020', 'Dokumenta datums. Jābūt vecākvienības datumu robežās. Obligāts DOK rindām.'),
    'REĢ_NR': ('1-15/1', 'Reģistrācijas numurs (līdz 30 z.). Obligāts DOK rindām.'),
    'IZVEIDOŠANAS_DATUMS': ('15.01.2020', 'Obligāts DOK rindām.'),
    'NOSŪTĪŠANAS_DATUMS': ('16.01.2020', 'Obligāts DOK rindām.'),
    'NOSŪTĪTĀJA_REĢ_NR': ('4.1-2/1234', 'Nosūtītāja reģistrācijas numurs (līdz 30 z.).'),
    'LIETAS_NR': ('1-15', 'Lietas (nomenklatūras) numurs. Obligāts DOK rindām.'),
    'GRUPA': ('Iekšējs / Saņemts / Nosūtīts', 'Brīvs teksts līdz 30 z.'),
    'ATSLĒGVĀRDI': ('protokols,dome', 'Atdala ar komatu.'),
    'ANOTĀCIJA': ('teksts', 'Dokumenta anotācija (līdz 500 z.).'),
    'TEHNISKĀ_INFORMĀCIJA': ('teksts', 'Tehniskā informācija par dokumentu (līdz 500 z.).'),
    'IEROBEŽOJUMA_DATUMS': ('31.12.2100', 'Obligāts, ja DOK PIEEJAMĪBA = Ierobežota. Ja Vispārēja — jābūt tukšam.'),
    'IEROBEŽOJUMA_PIEZĪMES': ('teksts', 'Līdz 30 z.'),
    'LIETOŠANAS_NOSACĪJUMI': ('teksts', 'Līdz 30 z.'),
}


def style_data_sheet(ws, rows, header_row=1):
    """Header + data rows with per-type row tint and frozen header."""
    ws.append(COLUMNS)
    for cell in ws[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER

    for r in rows:
        ws.append([r[column] for column in COLUMNS])
        fill = GV_FILL if r['TIPS'] == 'GV' else DOK_FILL
        for cell in ws[ws.max_row]:
            cell.fill = fill
            cell.border = BORDER
            cell.alignment = Alignment(vertical='top', wrap_text=False)

    widths = {'TIPS': 6, 'SAITE': 8, 'NOSAUKUMS': 44, 'SATURS': 40, 'PIEZĪMES': 30,
              'SĒRIJAS_KODS': 13, 'DATUMS_NO': 12, 'DATUMS_LĪDZ': 12,
              'DATUMA_PRECIZITĀTE': 18, 'ANOTĀCIJA': 38}
    for index, column in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(column, 20)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)


def build_xlsx(path):
    wb = Workbook()

    # Sheet 1 — the data the importer reads.
    ws = wb.active
    ws.title = 'DATI'
    style_data_sheet(ws, MIXED_ROWS)

    # Sheet 2 — instructions (ignored by the importer: it reads DATI / first sheet).
    doc = wb.create_sheet('INSTRUKCIJA')
    doc['A1'] = 'Kā aizpildīt importa failu'
    doc['A1'].font = TITLE_FONT
    intro = [
        '',
        'EKSPERIMENTĀLA FUNKCIJA. Imports ir izstrādes stadijā un nav rādītājs pārējā rīka kvalitātei.',
        'Pēc importa rezultāts obligāti jāpārbauda. Atsaukšanas iespējas nav.',
        '',
        'Rīks lasa TIKAI lapu "DATI" (vai pirmo lapu). Šī lapa ir tikai pamācībai.',
        'Kolonnu secība nav svarīga — svarīgi ir virsraksti pirmajā rindā.',
        'Neatpazītas kolonnas tiek ignorētas ar brīdinājumu.',
        '',
        'Katrai rindai kolonnā TIPS jānorāda:',
        '    GV  = glabājamā vienība',
        '    DOK = dokuments — tikai tekstuālos elektroniskos uzskaites sarakstos',
        '',
        'GV numurus piešķir sistēma pati, faila secībā — failā GV numura kolonnas nav.',
        'Datnes (failus) šis imports nepievieno — tikai aprakstus.',
        '',
        'Trīs veidi, kā failu var lietot (skat. lapu "PARAUGI"):',
        '    A) tikai GV rindas — izveido tikai glabājamās vienības',
        '    B) tikai DOK rindas ar SAITE = "GV:<numurs>" — pievieno dokumentus jau esošām vienībām',
        '    C) GV un DOK rindas kopā — izveido vienības ar to dokumentiem',
        '',
    ]
    for line in intro:
        doc.append([line])

    doc.append([])
    header_index = doc.max_row + 1
    doc.append(['KOLONNA', 'GV', 'DOK', 'PIEMĒRS / ATĻAUTĀS VĒRTĪBAS', 'PASKAIDROJUMS'])
    for cell in doc[header_index]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for column in COLUMNS:
        example, explanation = COLUMN_DOCS[column]
        if column in ('TIPS', 'SAITE'):
            gv_mark, dok_mark = 'oblig.' if column == 'TIPS' else '–', \
                                'oblig.' if column == 'TIPS' else 'jā'
        else:
            in_gv = column in GV_COLUMNS
            in_dok = column in DOK_COLUMNS
            shared = column in ('NOSAUKUMS', 'VALODA', 'PIEZĪMES', 'PIEEJAMĪBA')
            gv_mark = 'oblig.' if column in REQUIRED_GV else ('jā' if in_gv or shared else '–')
            dok_mark = 'oblig.' if column in REQUIRED_DOK else ('jā' if in_dok or shared else '–')
        doc.append([column, gv_mark, dok_mark, example, explanation])
        for cell in doc[doc.max_row]:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = BORDER

    doc.column_dimensions['A'].width = 26
    doc.column_dimensions['B'].width = 8
    doc.column_dimensions['C'].width = 8
    doc.column_dimensions['D'].width = 46
    doc.column_dimensions['E'].width = 74
    doc.freeze_panes = doc.cell(row=header_index + 1, column=1)

    # Sheet 3 — the three usage modes, each as its own small table.
    ex = wb.create_sheet('PARAUGI')
    ex['A1'] = 'Trīs lietojuma veidi'
    ex['A1'].font = TITLE_FONT
    ex.append([])

    def mini_table(title, rows, columns):
        ex.append([title])
        ex.cell(row=ex.max_row, column=1).font = Font(bold=True)
        ex.append(columns)
        for cell in ex[ex.max_row]:
            if cell.value:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal='center', wrap_text=True)
        for r in rows:
            ex.append([r[column] for column in columns])
            fill = GV_FILL if r['TIPS'] == 'GV' else DOK_FILL
            for cell in ex[ex.max_row]:
                cell.fill = fill
                cell.border = BORDER
        ex.append([])

    mini_table('A) Tikai glabājamās vienības',
               ITEMS_ONLY_ROWS,
               ['TIPS', 'SĒRIJAS_KODS', 'NOSAUKUMS', 'DATUMS_NO', 'DATUMS_LĪDZ', 'VALODA'])
    mini_table('B) Tikai dokumenti esošām vienībām (SAITE = GV:<numurs>)',
               RECORDS_ONLY_ROWS,
               ['TIPS', 'SAITE', 'NOSAUKUMS', 'DATUMS', 'REĢ_NR', 'IZVEIDOŠANAS_DATUMS',
                'NOSŪTĪŠANAS_DATUMS', 'LIETAS_NR', 'VALODA'])
    mini_table('C) Vienības ar dokumentiem (SAITE saista DOK ar GV)',
               MIXED_ROWS[:4],
               ['TIPS', 'SAITE', 'SĒRIJAS_KODS', 'NOSAUKUMS', 'DATUMS_NO', 'DATUMS_LĪDZ',
                'DATUMS', 'REĢ_NR', 'VALODA'])

    for index in range(1, 12):
        ex.column_dimensions[get_column_letter(index)].width = 24
    ex.column_dimensions['A'].width = 8

    wb.save(path)
    print(f'  {path}  (3 lapas)')


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    print('Rakstu paraugfailus:')
    write_csv(os.path.join(OUT_DIR, 'imports_paraugs.csv'), MIXED_ROWS)
    write_csv(os.path.join(OUT_DIR, 'imports_tikai_vienibas.csv'), ITEMS_ONLY_ROWS)
    write_csv(os.path.join(OUT_DIR, 'imports_tikai_dokumenti.csv'), RECORDS_ONLY_ROWS)
    build_xlsx(os.path.join(OUT_DIR, 'imports_paraugs.xlsx'))


if __name__ == '__main__':
    main()
